#Source: https://stackoverflow.com/questions/56049519/attributeerror-module-pptx-chart-has-no-attribute-data
import pptx
import pandas as pd

df = pd.read_excel("data.xlsx")

overall_report = pptx.Presentation("pres.pptx")


pres_slide = overall_report.slides[1]

slide_chart = pres_slide.shapes[20].chart

#replace chart data with the data from the excel above
chart_data = pptx.chart.data.CategoryChartData()  
chart_data.categories = df["Question"].values.tolist()

df1 = df.iloc[:,1:6].copy()

for col_idx, col in enumerate(df1.columns):
    print(col_idx,col,df1.iloc[:, col_idx].values)
    chart_data.add_series(col,(df1.iloc[:, col_idx].values))

#update data
slide_chart.replace_data(chart_data)