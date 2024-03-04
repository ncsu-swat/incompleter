#Source: https://stackoverflow.com/questions/53591660/python-docx-attributeerror-windowspath-object-has-no-attribute-seek
from pathlib import Path
import docx
from docx.shared import Cm

filepath = r"C:\Users\Admin\Desktop\img"
document = docx.Document()

for file in Path(filepath).iterdir():
#    paragraph = document.add_paragraph(Path(file).resolve().stem)
    document.add_picture(Path(file).absolute(), width=Cm(15.0))

document.save('test.docx')