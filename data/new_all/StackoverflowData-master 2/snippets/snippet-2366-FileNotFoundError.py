#Source: https://stackoverflow.com/questions/48516372/getting-typeerror-during-creation-of-bar-chart-from-excel-using-openpyxl-referen
#!/usr/bin/python

from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.chart import (
 Reference,
 Series,
 BarChart
)

wb = load_workbook('revenue.xlsx')
wsheet = wb.get_sheet_by_name('sales')
print (wsheet)


data =   Reference(wsheet, (5, 2), (5, 10))
categs = Reference(wsheet, (3, 2), (3, 10))

chart = BarChart()
chart.add_data(data=data)
chart.set_categories(categs)

chart.legend = None
chart.y_axis.majorGridlines = None
chart.varyColors = True
chart.title = "Sales By name"

wsheet.add_chart(chart, "H15")    

wb.save("revenue.xlsx")

print ("end")